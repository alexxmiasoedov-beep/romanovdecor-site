#!/usr/bin/env python3
# PDF-версия сметы из Excel (документы/сметы/смета-*.xlsx) в фирменном стиле UNICORE.
# Стили ячеек — копия STY из kp-generator/build_kp_losika.py (эталон вёрстки КП).
#
#   python3 build_smeta_pdf.py смета-Миля-9.xlsx
#   SC=/tmp/kp node ../kp-generator/topdf.js smeta "смета-Миля-9"
#
# Многостраничный документ: каждая страница — блок ровно 1123 px.
import os, sys, datetime
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else 'смета-Миля-9.xlsx'
SC = os.environ.get('SC', '/tmp/kp')
META = dict(
    title='Смета на ремонт квартиры под ключ',
    number='Смета № 270926M',
    date='27 сентября 2026 г.',
    customer='Юлия и Павел',
    address='г. Минск, ул. Миля, 9',
    basis='дизайн-проект Interior AK Design (Хворостов А.О.), 31.07.2026, 49 листов',
    area='53,88 м² без балкона, 57,03 м² с балконом, высота 2,73 м',
    rows_first=0,      # строк спецификации на первой странице (0 — только сводка)
    rows_page=50,      # строк спецификации на странице продолжения
)

wb = load_workbook(os.path.join(HERE, SRC), data_only=True)
M = wb['Смета']; PAR = wb['Параметры']
RATE = PAR['B3'].value

def u(v, dec=2):
    if v is None or v == '': return ''
    s = f'{v:,.{dec}f}'.replace(',', ' ').replace('.', ',')
    return s
def u0(v): return u(v, 0)
def vol(v):
    if v is None: return ''
    return u(v, 2) if abs(v - round(v)) > 1e-9 else u(v, 0)

# ── разбор листа «Смета» ────────────────────────────────────────────────────
items = []      # ('sec', name) | ('row', code, name, unit, vol, wp, mp, w, m, t) | ('sub', w, m, t)
totals = {}
notes = []
r = 6
while r <= M.max_row:
    b, c = M.cell(r, 2).value, M.cell(r, 3).value
    if c is None and b is None:
        r += 1; continue
    if b and str(b)[:1].isalpha() and str(b)[1:].isdigit():
        items.append(('row', b, c, M.cell(r, 4).value, M.cell(r, 5).value, M.cell(r, 6).value,
                      M.cell(r, 7).value, M.cell(r, 8).value, M.cell(r, 9).value, M.cell(r, 10).value))
    elif c == 'Итого по разделу':
        items.append(('sub', M.cell(r, 8).value, M.cell(r, 9).value, M.cell(r, 10).value))
    elif c and c[0].isdigit() and '. ' in c[:4]:
        items.append(('sec', c))
    elif c and (c.startswith('Итого') or c.startswith('ИТОГО') or c.startswith('Запас') or c.startswith('Накладные')
                or c.startswith('Доставка') or c.startswith('Непредвиденные') or c.startswith('Справочно')):
        totals[c] = (M.cell(r, 8).value, M.cell(r, 9).value, M.cell(r, 10).value)
    elif c and len(c) > 60:
        notes.append(c)
    r += 1

sections = []   # (name, w, m, t, n_rows)
cur = None
for it in items:
    if it[0] == 'sec': cur = [it[1], 0, 0, 0, 0]
    elif it[0] == 'row': cur[4] += 1
    elif it[0] == 'sub': cur[1:4] = it[1:4]; sections.append(tuple(cur))

T_BASE = totals['Итого прямые затраты']; T_RES = totals['Запас на материалы']; T_OVH = totals['Накладные и прибыль подрядчика']
T_DEL = totals['Доставка и подъём материалов']; T_UNF = totals['Непредвиденные']; T_TOT = totals['ИТОГО К ОПЛАТЕ, BYN']
per_m2 = [v for k, v in totals.items() if 'на м²' in k][0][2]
per_m2_bare = [v for k, v in totals.items() if 'без приборов' in k][0][2]

# ── стили (эталон build_kp_losika.py) ──────────────────────────────────────
B = '#cbbf9f'; GRID = f'border:1px solid {B}'; NUMS = 'font-variant-numeric:tabular-nums'
CELL = f'padding:4px 6px;{GRID};vertical-align:middle;font-size:9px;line-height:1.25;height:18px'
THB = f'{GRID};font-size:9px;letter-spacing:2px;text-transform:uppercase;font-weight:700;padding:5px 6px;vertical-align:middle;text-align:center'
THS = f'{GRID};font-weight:500;letter-spacing:1px;font-size:8.5px;padding:4px 6px;vertical-align:middle;text-align:center'
STY = dict(
 th_top_mat=f'background:#5a7236;color:#fff;{THB}',
 th_top_work=f'background:#a07a32;color:#fff;{THB}',
 th_top_tot=f'background:#1f1b14;color:#b8965a;{THB}',
 th_top_name=f'background:#2c2c2c;color:#fff;{GRID};padding:5px 10px;text-align:center;font-weight:600;letter-spacing:2px;text-transform:uppercase;font-size:9px;vertical-align:middle',
 th_sub_mat=f'background:#3d4a2e;color:#cfe0a8;{THS}',
 th_sub_work=f'background:#4a3a1f;color:#e8c88a;{THS}',
 td_name=f'{CELL};padding-left:10px;background:#f5efe2',
 td_unit=f'{CELL};text-align:center;white-space:nowrap;background:#f5efe2;color:#666',
 td_vol=f'{CELL};{NUMS};text-align:right;white-space:nowrap;background:#f5efe2',
 td_mat=f'{CELL};{NUMS};text-align:right;white-space:nowrap;background:#f0f5e0',
 td_work=f'{CELL};{NUMS};text-align:right;white-space:nowrap;background:#fdeed2',
 td_tot=f'{CELL};{NUMS};text-align:right;white-space:nowrap;background:#f5e2bb;font-weight:700;color:#2c2c2c;border-left:3px solid #b8965a')
SUBB = f'{GRID};{NUMS};font-weight:700;font-size:9px;padding:4px 6px;text-align:right;white-space:nowrap;height:18px;vertical-align:middle'
S_NAME = f'background:#eae0cb;{GRID};font-weight:700;font-size:9px;color:#3a3a3a;padding:4px 10px;vertical-align:middle'
S_MAT = f'background:#e2ecc8;color:#3a4a1c;{SUBB}'
S_WORK = f'background:#f8ddb0;color:#6a4a14;{SUBB}'
S_TOT = f'background:#eae0cb;color:#7a5614;{SUBB};border-left:3px solid #b8965a'
SEC = f'background:#2c2c2c;color:#efece7;{GRID};font-size:8.5px;letter-spacing:2px;text-transform:uppercase;font-weight:600;padding:4px 10px;height:18px'

COLG = ('<colgroup><col style="width:36%"><col style="width:6%"><col style="width:8%"><col style="width:9%">'
        '<col style="width:10%"><col style="width:9%"><col style="width:10%"><col style="width:12%"></colgroup>')
THEAD = (f'<thead><tr><th rowspan="2" style="{STY["th_top_name"]}">Наименование работ и затрат</th>'
         f'<th rowspan="2" style="{STY["th_top_name"]}">Ед.</th><th rowspan="2" style="{STY["th_top_name"]}">Объём</th>'
         f'<th colspan="2" style="{STY["th_top_mat"]}">Материалы</th><th colspan="2" style="{STY["th_top_work"]}">Работы</th>'
         f'<th rowspan="2" style="{STY["th_top_tot"]}">Итого, руб.</th></tr>'
         f'<tr><th style="{STY["th_sub_mat"]}">Цена</th><th style="{STY["th_sub_mat"]}">Сумма</th>'
         f'<th style="{STY["th_sub_work"]}">Цена</th><th style="{STY["th_sub_work"]}">Сумма</th></tr></thead>')

def tr_item(it):
    _, code, name, unit, v, wp, mp, w, m, t = it
    return (f'<tr><td style="{STY["td_name"]}">{name}</td><td style="{STY["td_unit"]}">{unit}</td>'
            f'<td style="{STY["td_vol"]}">{vol(v)}</td><td style="{STY["td_mat"]}">{u(mp) if mp else "—"}</td>'
            f'<td style="{STY["td_mat"]}">{u0(m) if m else "—"}</td><td style="{STY["td_work"]}">{u(wp) if wp else "—"}</td>'
            f'<td style="{STY["td_work"]}">{u0(w) if w else "—"}</td><td style="{STY["td_tot"]}">{u0(t)}</td></tr>')
def tr_sec(name): return f'<tr><td colspan="8" style="{SEC}">{name}</td></tr>'
def tr_sub(w, m, t):
    return (f'<tr><td colspan="3" style="{S_NAME}">Итого по разделу</td><td style="{S_MAT}"></td><td style="{S_MAT}">{u0(m)}</td>'
            f'<td style="{S_WORK}"></td><td style="{S_WORK}">{u0(w)}</td><td style="{S_TOT}">{u0(t)}</td></tr>')

logo = ('<svg width="44" height="44" viewBox="0 0 32 32" xmlns="http://www.w3.org/2000/svg">'
        '<rect width="32" height="32" rx="6" fill="#2c2c2c"/>'
        '<path d="M 9.6 8 L 9.6 15.5 A 6.4 6.4 0 0 0 22.4 15.5 L 22.4 8" fill="none" stroke="#b8965a" stroke-width="4"/></svg>')

def header(sub):
    return (f'<div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:5px;border-bottom:2px solid #b8965a;padding-bottom:5px">'
            f'<div style="display:flex;align-items:center;gap:10px">{logo}<div><div style="font-size:16px;letter-spacing:3px;font-weight:600;line-height:1.1">UNI<span style="color:#b8965a">CORE</span></div>'
            f'<div style="font-size:9px;color:#888;margin-top:3px;letter-spacing:1px;text-transform:uppercase">ремонт и микроцемент</div></div></div>'
            f'<div style="text-align:right;font-size:10px;color:#666;letter-spacing:1px"><div>{META["date"]}</div>'
            f'<div style="margin-top:2px;font-weight:600;color:#2c2c2c">{META["number"]}{sub}</div></div></div>')
def footer(p, n):
    return (f'<div style="margin-top:auto;padding-top:8px;border-top:1px solid #eee;display:flex;justify-content:space-between;align-items:center;font-size:9px;color:#888;letter-spacing:1px">'
            f'<div>romanovdecor.by · info@romanovdecor.by · +375 (33) 628-04-86 · Минск</div>'
            f'<div><span style="color:#b8965a;font-weight:600">Спасибо за доверие!</span> &nbsp;·&nbsp; стр. {p} / {n}</div></div>')
PAGE = 'width:100%;height:1123px;overflow:hidden;padding:14px 32px 12px;background:#fff;color:#2c2c2c;font-family:Geologica,Arial,sans-serif;display:flex;flex-direction:column;page-break-after:always'

# ── страница 1: сводка ──────────────────────────────────────────────────────
sum_rows = ''
for name, w, m, t, n in sections:
    sum_rows += (f'<tr><td style="{STY["td_name"]}">{name}</td><td style="{STY["td_unit"]}">{n} поз.</td>'
                 f'<td style="{STY["td_mat"]}">{u0(m)}</td><td style="{STY["td_work"]}">{u0(w)}</td><td style="{STY["td_tot"]}">{u0(t)}</td></tr>')
def sum_line(name, w, m, t, strong=False):
    st = 'font-weight:700;' if strong else ''
    return (f'<tr><td colspan="2" style="{S_NAME};{st}">{name}</td><td style="{S_MAT}">{u0(m) if m else ""}</td>'
            f'<td style="{S_WORK}">{u0(w) if w else ""}</td><td style="{S_TOT}">{u0(t)}</td></tr>')
sum_rows += sum_line('Итого прямые затраты', *T_BASE)
sum_rows += sum_line(f'Запас на материалы {PAR["B4"].value*100:.0f} %', None, T_RES[1], T_RES[2])
sum_rows += sum_line(f'Накладные и прибыль подрядчика {PAR["B7"].value*100:.0f} % от работ', T_OVH[0], None, T_OVH[2])
sum_rows += sum_line('Доставка и подъём материалов', None, T_DEL[1], T_DEL[2])
sum_rows += sum_line(f'Непредвиденные {PAR["B6"].value*100:.0f} %', None, None, T_UNF[2])

not_incl = notes[0] if notes else ''
cond = [n for n in notes[1:]]
cond_html = ''.join(f'<div style="margin-top:3px">• {n}</div>' for n in cond[:7])

page1 = f'''<div style="{PAGE}">
{header('')}
<div style="font-size:16px;font-weight:700;letter-spacing:2px;margin:9px 0 7px">{META['title']}</div>
<div style="display:flex;gap:10px;margin-bottom:6px">
  <div style="flex:1;padding:6px 12px;border-radius:6px;font-size:10px;line-height:1.42;background:#faf6ef;border:1px solid #ece2d2">
    <div style="font-size:8px;text-transform:uppercase;letter-spacing:2px;color:#b8965a;font-weight:600;margin-bottom:2px">Заказчик</div>
    <div><b>Заказчик:</b> {META['customer']}</div>
    <div><b>Адрес объекта:</b> {META['address']}</div>
    <div><b>Основание:</b> {META['basis']}</div>
    <div><b>Площадь:</b> {META['area']}</div>
  </div>
  <div style="flex:1;padding:6px 12px;border-radius:6px;font-size:10px;line-height:1.42;background:#2c2c2c;color:#efece7">
    <div style="font-size:8px;text-transform:uppercase;letter-spacing:2px;color:#b8965a;font-weight:600;margin-bottom:2px">Исполнитель</div>
    <div><b>UNICORE</b></div>
    <div>Алексей — менеджер проекта</div>
    <div>+375 (33) 628-04-86</div>
    <div style="color:#a09a92">info@romanovdecor.by · romanovdecor.by</div>
  </div>
</div>

<h3 style="font-size:10px;letter-spacing:3px;text-transform:uppercase;margin:9px 0 4px;color:#b8965a;font-weight:700">Сводка по разделам</h3>
<table style="width:100%;border-collapse:collapse;font-size:10px">
<colgroup><col style="width:46%"><col style="width:10%"><col style="width:14%"><col style="width:14%"><col style="width:16%"></colgroup>
<thead><tr><th style="{STY['th_top_name']}">Раздел</th><th style="{STY['th_top_name']}">Позиций</th>
<th style="{STY['th_top_mat']}">Материалы, руб.</th><th style="{STY['th_top_work']}">Работы, руб.</th><th style="{STY['th_top_tot']}">Итого, руб.</th></tr></thead>
<tbody>{sum_rows}</tbody></table>

<div style="text-align:right;margin:8px 0 0;letter-spacing:1px"><span style="font-size:11px;font-weight:600">ИТОГО К ОПЛАТЕ:</span> <span style="font-size:17px;font-weight:700;color:#b8965a">{u0(T_TOT[2])} руб.</span></div>
<div style="text-align:right;font-size:9px;color:#888;letter-spacing:1px;margin:3px 0 0">≈ €{u0(T_TOT[2]/RATE)} по курсу НБ РБ {u(RATE, 4)} BYN/EUR · {u0(per_m2)} руб. за м² всё включено · {u0(per_m2_bare)} руб. за м² без приборов и изделий раздела 12</div>

<div style="display:grid;grid-template-columns:repeat(3,1fr);gap:10px;margin:12px 0 0">
  <div style="padding:10px 14px;background:#f5efe2;border:1px solid #e3d7bd;border-radius:5px">
    <div style="color:#8a7346;letter-spacing:2px;text-transform:uppercase;font-size:8px;font-weight:600">Материалы</div>
    <div style="font-size:17px;font-weight:700;color:#2c2c2c;margin-top:3px">{u0(T_TOT[1])}</div></div>
  <div style="padding:10px 14px;background:#f5efe2;border:1px solid #e3d7bd;border-radius:5px">
    <div style="color:#8a7346;letter-spacing:2px;text-transform:uppercase;font-size:8px;font-weight:600">Работы</div>
    <div style="font-size:17px;font-weight:700;color:#2c2c2c;margin-top:3px">{u0(T_TOT[0])}</div></div>
  <div style="padding:10px 14px;background:#2c2c2c;border-radius:5px">
    <div style="color:#b8965a;letter-spacing:2px;text-transform:uppercase;font-size:8px;font-weight:600">Всего с непредвиденными</div>
    <div style="font-size:18px;font-weight:700;color:#b8965a;margin-top:3px">{u0(T_TOT[2])}</div></div>
</div>

<div style="margin-top:9px;padding:6px 12px;background:#f7f5f0;border-left:2px solid #b8965a;border-radius:3px;font-size:8.2px;color:#666;line-height:1.42">
<b style="color:#2c2c2c">Не включено.</b> {not_incl.replace('НЕ ВКЛЮЧЕНО (', '(').replace('НЕ ВКЛЮЧЕНО: ', '')}<br>
<b style="color:#2c2c2c">Условия.</b>{cond_html}
<div style="margin-top:3px">• Цены в белорусских рублях без НДС, срез рынка Минска на {META['date'][:-3]}; евро — справочно как валютная оговорка (ст. 298 ГК). Срок действия сметы — 14 дней.</div>
</div>
{footer(1, 'N')}
</div>'''

# ── страницы спецификации ───────────────────────────────────────────────────
body_rows = [(it, (tr_sec(it[1]) if it[0] == 'sec' else tr_item(it) if it[0] == 'row' else tr_sub(*it[1:]))) for it in items]
chunks, cur, h = [], [], 0
for it, html_row in body_rows:
    cost = 1
    if it[0] == 'row' and len(str(it[2])) > 62: cost = 2   # длинное имя — две строки
    if h + cost > META['rows_page'] and cur:
        chunks.append(cur); cur, h = [], 0
    if it[0] == 'sec' and h + 3 > META['rows_page']:        # заголовок не оставляем внизу
        chunks.append(cur); cur, h = [], 0
    cur.append(html_row); h += cost
if cur: chunks.append(cur)

pages = [page1]
for i, ch in enumerate(chunks):
    pages.append(f'''<div style="{PAGE}">
{header(' · спецификация')}
<h3 style="font-size:10px;letter-spacing:3px;text-transform:uppercase;margin:9px 0 4px;color:#b8965a;font-weight:700">Спецификация работ и материалов{' (продолжение)' if i else ''}</h3>
<table style="width:100%;border-collapse:collapse;font-size:9px">{COLG}{THEAD}<tbody>{''.join(ch)}</tbody></table>
{footer(i + 2, 'N')}
</div>''')
N = len(pages)
html = ('<!DOCTYPE html><html lang="ru"><head><meta charset="UTF-8">'
        '<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>'
        '<link href="https://fonts.googleapis.com/css2?family=Geologica:wght@300;400;500;600;700&display=swap" rel="stylesheet">'
        '<style>@page{size:A4;margin:0}*{margin:0;padding:0;box-sizing:border-box}body{width:210mm}</style></head><body>'
        + ''.join(pages).replace('/ N<', f'/ {N}<') + '</body></html>')
os.makedirs(SC + '/kp', exist_ok=True)
open(SC + '/kp/smeta.html', 'w').write(html)
print(f'страниц {N}; итог {u0(T_TOT[2])} руб.; html → {SC}/kp/smeta.html')
