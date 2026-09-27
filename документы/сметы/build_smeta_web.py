#!/usr/bin/env python3
# Веб-версия сметы (HTML для публикации как артефакт) из Excel-сметы.
#   python3 build_smeta_web.py смета-Миля-9.xlsx /путь/вывод.html
import os, sys, html as H
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else 'смета-Миля-9.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, 'смета-Миля-9.html')
META = dict(title='Смета Миля 9', number='Смета № 270926M', date='27 сентября 2026',
            customer='Юлия и Павел', address='г. Минск, ул. Миля, 9',
            basis='дизайн-проект Interior AK Design (Хворостов А.О.), 31.07.2026, 49 листов',
            area='53,88 м² без балкона, 57,03 м² с балконом, высота 2,73 м')

wb = load_workbook(os.path.join(HERE, SRC), data_only=True)
M = wb['Смета']; PAR = wb['Параметры']; RATE = PAR['B3'].value
def u(v, d=2):
    return '' if v in (None, '') else f'{v:,.{d}f}'.replace(',', ' ').replace('.', ',')
def u0(v): return u(v, 0)
def vol(v): return '' if v is None else (u(v, 2) if abs(v - round(v)) > 1e-9 else u(v, 0))
e = H.escape

items, totals, notes = [], {}, []
r = 6
while r <= M.max_row:
    b, c = M.cell(r, 2).value, M.cell(r, 3).value
    if b and str(b).startswith('M'):
        items.append(('row', b, c, M.cell(r, 4).value, M.cell(r, 5).value, M.cell(r, 6).value,
                      M.cell(r, 7).value, M.cell(r, 8).value, M.cell(r, 9).value, M.cell(r, 10).value))
    elif c == 'Итого по разделу':
        items.append(('sub', M.cell(r, 8).value, M.cell(r, 9).value, M.cell(r, 10).value))
    elif c and c[0].isdigit() and '. ' in c[:4]:
        items.append(('sec', c))
    elif c and (c.startswith(('Итого', 'ИТОГО', 'Запас', 'Накладные', 'Доставка', 'Непредвиденные', 'Справочно'))):
        totals[c] = (M.cell(r, 8).value, M.cell(r, 9).value, M.cell(r, 10).value)
    elif c and len(c) > 60:
        notes.append(c)
    r += 1
sections, cur = [], None
for it in items:
    if it[0] == 'sec': cur = [it[1], 0, 0, 0, 0]
    elif it[0] == 'row': cur[4] += 1
    elif it[0] == 'sub': cur[1:4] = it[1:4]; sections.append(tuple(cur))
T_BASE = totals['Итого прямые затраты']; T_RES = totals['Запас на материалы']; T_OVH = totals['Накладные и прибыль подрядчика']
T_DEL = totals['Доставка и подъём материалов']; T_UNF = totals['Непредвиденные']; T_TOT = totals['ИТОГО К ОПЛАТЕ, BYN']
per_m2 = [v for k, v in totals.items() if 'на м²' in k][0][2]
work_m2 = [v for k, v in totals.items() if 'только работы' in k][0][2]

CSS = '''
:root{--bg:#f5efe2;--paper:#ffffff;--ink:#2c2c2c;--ink-soft:#5a5651;--muted:#8a8378;--gold:#b8965a;--gold-ink:#8a7346;--rule:#cbbf9f;--rule-soft:#e8dfcc;
 --card:#faf6ef;--dark:#2c2c2c;--dark-ink:#efece7;--dark-muted:#a09a92;
 --mat-head:#5a7236;--mat-bg:#f0f5e0;--mat-sub:#e2ecc8;--mat-ink:#3a4a1c;
 --work-head:#a07a32;--work-bg:#fdeed2;--work-sub:#f8ddb0;--work-ink:#6a4a14;
 --tot-bg:#f5e2bb;--tot-sub:#eae0cb;--tot-ink:#7a5614;--sec-bg:#2c2c2c;--sec-ink:#efece7;--focus:#b8965a}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){color-scheme:dark;--bg:#181613;--paper:#221f1b;--ink:#efece7;--ink-soft:#cfc9bf;--muted:#a09a92;--gold:#cfae72;--gold-ink:#cfae72;--rule:#4a4237;--rule-soft:#332e27;
 --card:#2a2622;--dark:#100f0d;--dark-ink:#efece7;--dark-muted:#a09a92;
 --mat-head:#4a5f2c;--mat-bg:#262d1c;--mat-sub:#2f3a22;--mat-ink:#cfe0a8;
 --work-head:#8a6828;--work-bg:#32281a;--work-sub:#3e3120;--work-ink:#e8c88a;
 --tot-bg:#3b3120;--tot-sub:#2f2a22;--tot-ink:#e0c48c;--sec-bg:#0f0e0c;--sec-ink:#efece7}}
:root[data-theme="dark"]{color-scheme:dark;--bg:#181613;--paper:#221f1b;--ink:#efece7;--ink-soft:#cfc9bf;--muted:#a09a92;--gold:#cfae72;--gold-ink:#cfae72;--rule:#4a4237;--rule-soft:#332e27;
 --card:#2a2622;--dark:#100f0d;--dark-ink:#efece7;--dark-muted:#a09a92;
 --mat-head:#4a5f2c;--mat-bg:#262d1c;--mat-sub:#2f3a22;--mat-ink:#cfe0a8;
 --work-head:#8a6828;--work-bg:#32281a;--work-sub:#3e3120;--work-ink:#e8c88a;
 --tot-bg:#3b3120;--tot-sub:#2f2a22;--tot-ink:#e0c48c;--sec-bg:#0f0e0c;--sec-ink:#efece7}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font-family:Geologica,Arial,sans-serif;font-size:14px;line-height:1.45;padding-inline:16px;padding-block:20px 40px}
.doc{max-width:980px;margin:0 auto;background:var(--paper);border:1px solid var(--rule-soft);border-radius:8px;padding:clamp(16px,3vw,32px)}
.head{display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap;border-bottom:2px solid var(--gold);padding-bottom:10px}
.brand{display:flex;align-items:center;gap:10px}
.brand .n{font-size:18px;letter-spacing:3px;font-weight:600;line-height:1.1}.brand .n b{color:var(--gold);font-weight:600}
.brand .s{font-size:9px;color:var(--muted);letter-spacing:1.6px;text-transform:uppercase;margin-top:3px}
.meta{text-align:right;font-size:12px;color:var(--muted);letter-spacing:.5px}.meta b{display:block;color:var(--ink);font-weight:600}
h1{font-size:clamp(18px,3vw,24px);letter-spacing:1.5px;font-weight:700;margin:18px 0 12px;text-wrap:balance}
h2{font-size:11px;letter-spacing:3px;text-transform:uppercase;color:var(--gold-ink);font-weight:700;margin:26px 0 8px}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));gap:10px}
.card{padding:10px 14px;border-radius:6px;font-size:13px;line-height:1.5;background:var(--card);border:1px solid var(--rule-soft)}
.card.dark{background:var(--dark);color:var(--dark-ink);border-color:transparent}.card.dark .dim{color:var(--dark-muted)}
.eyebrow{font-size:9px;text-transform:uppercase;letter-spacing:2px;color:var(--gold);font-weight:600;margin-bottom:4px}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin-top:14px}
.kpi{padding:12px 14px;background:var(--card);border:1px solid var(--rule-soft);border-radius:6px}
.kpi .l{font-size:9px;letter-spacing:2px;text-transform:uppercase;color:var(--gold-ink);font-weight:600}
.kpi .v{font-size:22px;font-weight:700;margin-top:4px;font-variant-numeric:tabular-nums}
.kpi.total{background:var(--dark);border-color:transparent}.kpi.total .l,.kpi.total .v{color:var(--gold)}
.total-line{margin-top:16px;display:flex;justify-content:flex-end;align-items:baseline;gap:10px;flex-wrap:wrap;letter-spacing:1px}
.total-line .lbl{font-size:12px;font-weight:600}.total-line .sum{font-size:26px;font-weight:700;color:var(--gold);font-variant-numeric:tabular-nums}
.sub{text-align:right;font-size:12px;color:var(--muted);margin-top:4px}
.wrap{overflow-x:auto;-webkit-overflow-scrolling:touch;border:1px solid var(--rule);border-radius:4px}
table{width:100%;border-collapse:collapse;font-size:12.5px;min-width:640px}
th,td{border:1px solid var(--rule);padding:6px 8px;vertical-align:middle}
th{font-size:9.5px;letter-spacing:1.6px;text-transform:uppercase;font-weight:700;text-align:center}
th.name{background:var(--sec-bg);color:var(--sec-ink)}th.mat{background:var(--mat-head);color:#fff}th.work{background:var(--work-head);color:#fff}th.tot{background:#1f1b14;color:#b8965a}
th.mat.s{background:var(--mat-sub);color:var(--mat-ink);font-weight:500;letter-spacing:1px}th.work.s{background:var(--work-sub);color:var(--work-ink);font-weight:500;letter-spacing:1px}
td{font-variant-numeric:tabular-nums}
td.name{text-align:left}td.unit{text-align:center;color:var(--muted);white-space:nowrap;font-size:11.5px}
td.num{text-align:right;white-space:nowrap}td.mat{background:var(--mat-bg)}td.work{background:var(--work-bg)}
td.tot{background:var(--tot-bg);font-weight:700;border-left:3px solid var(--gold)}
tr.sec td{background:var(--sec-bg);color:var(--sec-ink);font-size:10.5px;letter-spacing:2px;text-transform:uppercase;font-weight:600;padding:7px 10px}
tr.subt td{font-weight:700;background:var(--tot-sub)}tr.subt td.mat{background:var(--mat-sub);color:var(--mat-ink)}tr.subt td.work{background:var(--work-sub);color:var(--work-ink)}tr.subt td.tot{color:var(--tot-ink)}
tr.line td{background:var(--tot-sub);font-weight:600}
.note{margin-top:14px;padding:10px 14px;background:var(--card);border-left:3px solid var(--gold);border-radius:3px;font-size:12px;color:var(--ink-soft);line-height:1.5}
.note b{color:var(--ink)}.note p{margin:4px 0}
.foot{margin-top:26px;padding-top:10px;border-top:1px solid var(--rule-soft);display:flex;justify-content:space-between;gap:10px;flex-wrap:wrap;font-size:11px;color:var(--muted);letter-spacing:.5px}
.foot .t{color:var(--gold);font-weight:600}
a{color:var(--gold-ink)}:focus-visible{outline:2px solid var(--focus);outline-offset:2px}
@media (prefers-reduced-motion: reduce){*{scroll-behavior:auto}}
'''
logo = ('<svg width="44" height="44" viewBox="0 0 32 32" aria-hidden="true"><rect width="32" height="32" rx="6" fill="#2c2c2c"/>'
        '<path d="M 9.6 8 L 9.6 15.5 A 6.4 6.4 0 0 0 22.4 15.5 L 22.4 8" fill="none" stroke="#b8965a" stroke-width="4"/></svg>')

sum_rows = ''.join(f'<tr><td class="name">{e(n)}</td><td class="unit">{k} поз.</td><td class="num mat">{u0(m)}</td><td class="num work">{u0(w)}</td><td class="num tot">{u0(t)}</td></tr>'
                   for n, w, m, t, k in sections)
def line(name, w, m, t):
    return f'<tr class="line"><td class="name" colspan="2">{e(name)}</td><td class="num mat">{u0(m) if m else ""}</td><td class="num work">{u0(w) if w else ""}</td><td class="num tot">{u0(t)}</td></tr>'
sum_rows += line('Итого прямые затраты', *T_BASE)
sum_rows += line(f'Запас на материалы {PAR["B4"].value*100:.0f} %', None, T_RES[1], T_RES[2])
sum_rows += line(f'Накладные и прибыль подрядчика {PAR["B7"].value*100:.0f} % от работ', T_OVH[0], None, T_OVH[2])
sum_rows += line('Доставка и подъём материалов', None, T_DEL[1], T_DEL[2])
sum_rows += line(f'Непредвиденные {PAR["B6"].value*100:.0f} %', None, None, T_UNF[2])

spec = ''
for it in items:
    if it[0] == 'sec': spec += f'<tr class="sec"><td colspan="8">{e(it[1])}</td></tr>'
    elif it[0] == 'sub': spec += f'<tr class="subt"><td class="name" colspan="3">Итого по разделу</td><td class="mat"></td><td class="num mat">{u0(it[2])}</td><td class="work"></td><td class="num work">{u0(it[1])}</td><td class="num tot">{u0(it[3])}</td></tr>'
    else:
        _, code, name, unit, v, wp, mp, w, m, t = it
        spec += (f'<tr><td class="name">{e(name)}</td><td class="unit">{e(unit or "")}</td><td class="num">{vol(v)}</td>'
                 f'<td class="num mat">{u(mp) if mp else "—"}</td><td class="num mat">{u0(m) if m else "—"}</td>'
                 f'<td class="num work">{u(wp) if wp else "—"}</td><td class="num work">{u0(w) if w else "—"}</td><td class="num tot">{u0(t)}</td></tr>')

not_incl = notes[0].replace('НЕ ВКЛЮЧЕНО (', '(') if notes else ''
cond = ''.join(f'<p>• {e(n)}</p>' for n in notes[1:])

page = f'''<title>{e(META["title"])}</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Geologica:wght@400;500;600;700&display=swap">
<style>{CSS}</style>
<div class="doc">
<div class="head"><div class="brand">{logo}<div><div class="n">UNI<b>CORE</b></div><div class="s">ремонт и микроцемент</div></div></div>
<div class="meta">{META["date"]} г.<b>{META["number"]}</b></div></div>
<h1>Смета на ремонт квартиры под ключ</h1>
<div class="cards">
<div class="card"><div class="eyebrow">Заказчик</div><div><b>Заказчик:</b> {META["customer"]}</div><div><b>Адрес объекта:</b> {META["address"]}</div><div><b>Основание:</b> {META["basis"]}</div><div><b>Площадь:</b> {META["area"]}</div></div>
<div class="card dark"><div class="eyebrow">Исполнитель</div><div><b>UNICORE</b></div><div>Алексей — менеджер проекта</div><div>+375 (33) 628-04-86</div><div class="dim">info@romanovdecor.by · romanovdecor.by</div></div>
</div>
<div class="kpis">
<div class="kpi"><div class="l">Материалы</div><div class="v">{u0(T_TOT[1])}</div></div>
<div class="kpi"><div class="l">Работы</div><div class="v">{u0(T_TOT[0])}</div></div>
<div class="kpi"><div class="l">За м² без балкона</div><div class="v">{u0(per_m2)}</div></div>
<div class="kpi total"><div class="l">Итого, руб.</div><div class="v">{u0(T_TOT[2])}</div></div>
</div>
<div class="sub">≈ €{u0(T_TOT[2]/RATE)} по курсу НБ РБ {u(RATE,4)} BYN/EUR · работы {u0(work_m2)} руб. за м² · цены без НДС</div>

<h2>Сводка по разделам</h2>
<div class="wrap"><table>
<thead><tr><th class="name">Раздел</th><th class="name">Позиций</th><th class="mat">Материалы, руб.</th><th class="work">Работы, руб.</th><th class="tot">Итого, руб.</th></tr></thead>
<tbody>{sum_rows}</tbody></table></div>
<div class="total-line"><span class="lbl">ИТОГО К ОПЛАТЕ:</span><span class="sum">{u0(T_TOT[2])} руб.</span></div>

<h2>Спецификация работ и материалов</h2>
<div class="wrap"><table>
<thead><tr><th class="name" rowspan="2">Наименование работ и затрат</th><th class="name" rowspan="2">Ед.</th><th class="name" rowspan="2">Объём</th><th class="mat" colspan="2">Материалы</th><th class="work" colspan="2">Работы</th><th class="tot" rowspan="2">Итого, руб.</th></tr>
<tr><th class="mat s">Цена</th><th class="mat s">Сумма</th><th class="work s">Цена</th><th class="work s">Сумма</th></tr></thead>
<tbody>{spec}</tbody></table></div>

<div class="note"><p><b>Не включено.</b> {e(not_incl)}</p><p><b>Условия.</b></p>{cond}
<p>• Цены в белорусских рублях без НДС, срез рынка Минска на {META["date"]}; евро — справочно как валютная оговорка (ст. 298 ГК). Срок действия сметы — 14 дней.</p></div>
<div class="foot"><div>romanovdecor.by · info@romanovdecor.by · +375 (33) 628-04-86 · Минск</div><div class="t">Спасибо за доверие!</div></div>
</div>
'''
open(OUT, 'w', encoding='utf-8').write(page)
print('written', OUT, len(page), 'bytes')
